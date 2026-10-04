"""Traceability matrix, standards coverage, gap analysis and exports."""
import csv
import json

from openpyxl import load_workbook

from aegis_rmf.pipeline import export_outputs
from aegis_rmf.traceability import exporters


def row_for(result, risk_id):
    return next(row for row in result.matrix if row.risk_id == risk_id)


def test_matrix_has_one_row_per_risk_and_agrees_with_the_records(result, records):
    assert [row.risk_id for row in result.matrix] == sorted(records)
    for row in result.matrix:
        record = records[row.risk_id]
        assert (row.rpn_pre, row.rpn_post, row.region_post) == (record.pre.rpn, record.post.rpn, record.post.region.value)
        assert row.component_id == record.component_id and row.harm_id == record.harm_id
        assert sorted(row.control_ids) == sorted(c.control_id for c in record.controls)


def test_row_carries_full_cross_references(result):
    row = row_for(result, "RISK_SW_02")
    assert row.component_name == "Dose Calculation Software" and row.subsystem == "Therapy software"
    assert row.requirement_ids == ["REQ_SW_004", "REQ_SW_005"]
    assert row.verification_refs == ["VER_SW_004 PASS TR_0132", "VER_SW_005 FAIL TR_0133"]
    assert row.design_refs[0] == "SRS_003 section 4.1 (DFC_SW_02)"
    assert "REQ_SW_005 at SRS_003 section 4.1" in row.design_refs
    assert {"ISO 14971 cl 5.4", "ISO 14971 cl 7.2", "IEC 62304 cl 7.2", "IEC 62304 cl 7.4"} <= set(row.standard_refs)
    assert "verification FAIL" in row.controls[0] and "credited P0" in row.controls[0]


def test_row_reports_missing_verification_record(result):
    assert "REQ_LBL_002 has no verification record" in row_for(result, "RISK_RSV_01").verification_refs


def test_every_link_is_bidirectional(result):
    assert result.bidirectional["bidirectional"] is True
    assert result.bidirectional["links_checked"] == 113 and result.bidirectional["broken_links"] == []


def test_gap_report_finds_the_seeded_problems(result):
    found = {(gap.rule.split()[0], gap.subject_id) for gap in result.gaps}
    assert {("G1", "CMP_ACC_01"), ("G2", "RISK_SNS_03"), ("G3", "RCM_IFU_02"), ("G4", "RCM_SW_03"), ("G4", "RCM_COM_02"),
            ("G6", "RISK_SW_02"), ("G8", "RISK_INF_01"), ("G8", "RISK_SW_05")} <= found
    assert [gap.level for gap in result.gaps] == sorted((gap.level for gap in result.gaps), key=["BLOCKER", "MAJOR", "ACTION"].index)
    belt_clip = next(gap for gap in result.gaps if gap.subject_id == "CMP_ACC_01")
    assert "similar incident reports" in belt_clip.message


def test_no_gap_is_raised_for_a_clean_risk(result):
    subjects = {gap.subject_id for gap in result.gaps}
    assert "RISK_PWR_02" not in subjects and "RCM_PWR_03" not in subjects


def test_standards_coverage(result):
    coverage = {clause.key: clause for clause in result.coverage}
    assert len(coverage) == len(result.clauses)
    assert coverage["ISO 14971 cl 5.4"].risk_items == 39 and coverage["ISO 14971 cl 5.4"].status == "Evidenced"
    assert coverage["IEC 62304 cl 7.2"].requirements > 10
    assert "RMP_006 section 5" in coverage["ISO 14971 cl 7.5"].document_sections
    assert coverage["ISO 14971 cl 4.3"].status == "No evidence in file"


def test_exports(result, tmp_path):
    paths = export_outputs(result, tmp_path)
    with paths["matrix_csv"].open(encoding="utf8") as handle:
        rows = list(csv.DictReader(handle))
    assert len(rows) == 39 and len(rows[0]) == len(exporters.MATRIX_COLUMNS)
    assert rows[0]["Risk ID"] == "RISK_ALM_01" and rows[0]["Design specification references"]

    workbook = load_workbook(paths["matrix_xlsx"])
    assert workbook.sheetnames == ["Traceability Matrix", "Standards Coverage", "Gap Report"]
    assert workbook["Traceability Matrix"].max_row == 40 and workbook["Gap Report"].max_row == len(result.gaps) + 1

    page = paths["matrix_html"].read_text(encoding="utf8")
    assert page.count("<tr>") == 40 and 'class="unacceptable"' in page

    register = json.loads(paths["risk_register"].read_text(encoding="utf8"))
    assert len(register) == 39 and register[0]["audit"][0]["tool"] == "resolve_component"

    summary = paths["summary"].read_text(encoding="utf8")
    assert "RISK_SW_02" in summary and "Broken links: 0" in summary

    report = json.loads(paths["evaluation"].read_text(encoding="utf8"))
    assert report["ragas"]["status"] == "skipped"
    manifest = json.loads(paths["manifest"].read_text(encoding="utf8"))
    assert manifest["summary"]["risks"] == 39 and manifest["graph"] == result.store.stats()
    assert "proposals" not in paths
