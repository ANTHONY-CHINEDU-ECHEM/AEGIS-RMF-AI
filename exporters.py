"""Export the traceability matrix, standards coverage and gap report to CSV, Excel, HTML and JSON."""
from __future__ import annotations

import csv
import html
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from aegis_rmf.domain import RiskRecord
from aegis_rmf.traceability.gaps import Gap
from aegis_rmf.traceability.matrix import ClauseCoverage, TraceRow

SEPARATOR = "; "
REGION_COLORS = {"ACCEPTABLE": "C6EFCE", "REVIEW": "FFEB9C", "UNACCEPTABLE": "FFC7CE"}

MATRIX_COLUMNS: list[tuple[str, str, int]] = [
    ("Risk ID", "risk_id", 13), ("Title", "title", 30), ("Component ID", "component_id", 13),
    ("Component", "component_name", 26), ("Subsystem", "subsystem", 14), ("Failure mode ID", "failure_mode_id", 12),
    ("Failure mode", "failure_mode", 36), ("Cause", "cause", 34), ("Hazard", "hazard", 22),
    ("Hazardous situation ID", "hazardous_situation_id", 12), ("Hazardous situation", "hazardous_situation", 40),
    ("Harm ID", "harm_id", 18), ("Harm", "harm", 32), ("Detection before control", "detection_method", 20),
    ("S pre", "severity_pre", 6), ("P pre", "probability_pre", 6), ("D pre", "detectability_pre", 6),
    ("RPN pre", "rpn_pre", 8), ("Region pre", "region_pre", 15), ("Priority pre", "priority_pre", 10),
    ("Probability basis", "probability_basis", 16), ("Similar incidents", "field_matched_count", 10),
    ("Risk control measures", "controls", 60), ("Requirements", "requirement_ids", 24),
    ("Verification evidence", "verification_refs", 34),
    ("S post", "severity_post", 6), ("P post", "probability_post", 6), ("D post", "detectability_post", 6),
    ("RPN post", "rpn_post", 8), ("Region post", "region_post", 15), ("Priority post", "priority_post", 10),
    ("RPN reduction pct", "rpn_reduction_pct", 10), ("Disposition", "disposition", 28),
    ("Design specification references", "design_refs", 46), ("Standard clause references", "standard_refs", 40),
    ("Incident evidence", "incident_refs", 30),
]
COVERAGE_COLUMNS = [("Standard", "designation", 16), ("Clause", "clause", 8), ("Topic", "topic", 50),
                    ("Risk items", "risk_items", 10), ("Requirements", "requirements", 12),
                    ("Document sections", "document_sections", 50), ("Status", "status", 20)]
GAP_COLUMNS = [("Level", "level", 10), ("Rule", "rule", 40), ("Subject", "subject_id", 14),
               ("Clause", "clause", 18), ("Finding", "message", 110)]


def _cell(value) -> str | int | float:
    return SEPARATOR.join(str(item) for item in value) if isinstance(value, list) else value


def _table(items: list, columns: list[tuple[str, str, int]]) -> tuple[list[str], list[list]]:
    return [header for header, _, _ in columns], [[_cell(getattr(item, key)) for _, key, _ in columns] for item in items]


def write_csv(path: Path, items: list, columns: list[tuple[str, str, int]]) -> Path:
    headers, rows = _table(items, columns)
    with Path(path).open("w", newline="", encoding="utf8") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows(rows)
    return path


def write_risk_register(path: Path, records: list[RiskRecord]) -> Path:
    payload = [record.model_dump(mode="json") for record in records]
    Path(path).write_text(json.dumps(payload, indent=2), encoding="utf8")
    return path


def _sheet(workbook: Workbook, title: str, items: list, columns: list[tuple[str, str, int]]):
    sheet = workbook.create_sheet(title)
    headers, rows = _table(items, columns)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    for index, (_, _, width) in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
        cell = sheet.cell(row=1, column=index)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.value in REGION_COLORS:
                cell.fill = PatternFill("solid", fgColor=REGION_COLORS[cell.value])
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
    return sheet


def write_xlsx(path: Path, matrix: list[TraceRow], coverage: list[ClauseCoverage], gaps: list[Gap]) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)
    _sheet(workbook, "Traceability Matrix", matrix, MATRIX_COLUMNS)
    _sheet(workbook, "Standards Coverage", coverage, COVERAGE_COLUMNS)
    _sheet(workbook, "Gap Report", gaps, GAP_COLUMNS)
    workbook.save(path)
    return path


def write_html(path: Path, title: str, matrix: list[TraceRow], generated: str) -> Path:
    headers, rows = _table(matrix, MATRIX_COLUMNS)
    head = "".join(f"<th>{html.escape(header)}</th>" for header in headers)
    body = []
    for row in rows:
        cells = []
        for value in row:
            style = f' class="{value.lower()}"' if value in REGION_COLORS else ""
            cells.append(f"<td{style}>{html.escape(str(value)).replace(SEPARATOR, '<br>')}</td>")
        body.append(f"<tr>{''.join(cells)}</tr>")
    css = (
        "body{font:13px Arial,sans-serif;margin:24px;color:#1b1f24}"
        "table{border-collapse:collapse}"
        "th{background:#1f3864;color:#fff;position:sticky;top:0;padding:6px;text-align:left}"
        "td{border:1px solid #c9ced6;padding:6px;vertical-align:top;min-width:70px}"
        ".acceptable{background:#c6efce}.review{background:#ffeb9c}.unacceptable{background:#ffc7ce}"
    )
    document = (
        f"<!DOCTYPE html><html lang=\"en\"><head><meta charset=\"utf8\"><title>{html.escape(title)}</title>"
        f"<style>{css}</style></head><body><h1>{html.escape(title)}</h1><p>{html.escape(generated)}</p>"
        f"<table><thead><tr>{head}</tr></thead><tbody>{''.join(body)}</tbody></table></body></html>"
    )
    Path(path).write_text(document, encoding="utf8")
    return path
